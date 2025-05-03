import os
from faker import Faker
from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen import canvas
import zipfile
import rarfile
from py7zr import SevenZipFile

fake = Faker()

def generate_docx(path):
    doc = Document()
    doc.add_paragraph(fake.text())
    doc.save(path)

def generate_xlsx(path):
    wb = Workbook()
    ws = wb.active
    for row in range(10):
        ws.append([fake.name(), fake.random_int(1, 100)])
    wb.save(path)

def generate_pdf(path):
    c = canvas.Canvas(path)
    c.drawString(50, 800, fake.text())
    c.save()

def create_archive(archive_path, files_to_pack):
    if archive_path.endswith(".zip"):
        with zipfile.ZipFile(archive_path, "w") as z:
            for file in files_to_pack:
                z.write(file)
    elif archive_path.endswith(".rar"):
        with rarfile.RarFile(archive_path, "w") as r:
            for file in files_to_pack:
                r.add(file)
    elif archive_path.endswith(".7z"):
        with SevenZipFile(archive_path, "w") as z:
            for file in files_to_pack:
                z.write(file)

def gen_files():
    # Создание структуры
    os.makedirs("storage", exist_ok=True)
    generate_docx("storage/document1.docx")
    generate_xlsx("storage/data.xlsx")
    generate_pdf("storage/report.pdf")
    print("Документы сгенерированы!")
    temp_files = ["storage/document2.docx", "storage/data2.xlsx", "storage/report2.pdf"]
    generate_docx(temp_files[0])
    generate_xlsx(temp_files[1])
    generate_pdf(temp_files[2])
    create_archive("storage/archive.zip", temp_files)
    print("Архив добавлен.")

